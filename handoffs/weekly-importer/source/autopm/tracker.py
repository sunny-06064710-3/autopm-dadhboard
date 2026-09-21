"""Offline, header-based reconciliation against a user's existing All Tracker."""

from collections import Counter, defaultdict
from datetime import date, datetime, time, timedelta
import hashlib
import json
from pathlib import Path
import warnings
from zipfile import ZipFile

from .normalize import clean_text, normalize_header, normalize_project_id, parse_date
from .workbook import WorkbookError, MAX_FILE_BYTES, MAX_UNCOMPRESSED_BYTES

# The workbook itself marks P:R as immutable after project kick-off.  Keep this
# rule positional because those three business columns are explicitly governed
# by the template, even if somebody later changes their header wording.
PROTECTED_TRACKER_COLUMNS = frozenset({16, 17, 18})

HEADERS = {
    "project_id": ["Project Number", "Project ID", "Project ID (Manual)"],
    "project_name": [
        "Project Name , Description",
        "Project Name",
        "Project name (Manual)",
    ],
    "sku": ["SKUs Kicked Off", "SKU", "SKUs", "Project SKU (Manual)"],
    "type": ["Project Type", "Project Type (Manual)"],
    "brand": ["Brand", "Brand (Manual)"],
    "category": ["Category"],
    "sub_category": ["Sub Category", "Sub Category (Manual)"],
    "region": ["Region", "Region (Manual)"],
    "status": ["Eng Status", "Project Status", "Project Status (Manual)"],
    "npi_lead": ["NPI Project Lead", "NPI Owner (Manual)"],
    "npd_lead": ["NPD Project Lead", "NPD Owner(Manual)"],
    "pmo": ["PMO", "PMO Owner (Manual)"],
    "quality": ["Quality Owner (Manual)"],
    "factory": ["Manufacturing Factory", "Factory", "Factory (Manual)"],
    "capacity": ["Capacity", "Capacity / Forecast (Manual)"],
    "tooling": ["Tooling", "Tooling (Manual)"],
    "current_progress": [
        "Current Progress",
        "Current Progress (Manual)",
        "Engineering Remarks [INITIAL MM,DD: Comment]",
    ],
    "update_this_week": ["Eng Update This Week", "Update This Week", "Update This Week (Manual)", "Current stage"],
    "loa_13weeks_plan": ["LOA-13weeks plan", "LOA-13weeks plan(Manual)", "LOA-13wks plan"],
    "jira_summary": ["Jira summary", "Jira summary (Manual)", "Jira summary （Manual）"],
    "jira_link": ["Jira link", "Jira link (Manual)"],
    "next_plm": ["Next PLM", "Next Action"],
    "kick_off_date": ["Eng , OEM Kick Off", "Kick Off Date", "Kick Off Date (Manual)"],
    "dqtp_finish_date": ["DQTP Finish Date", "DQTP Finish Date (Manual)"],
    "mp_aw_date": ["MP AW Date", "MP AW Date (Manual)"],
    "compliance_complete_date": [
        "Compliance Complete Date",
        "Compliance Complete Date (Manual)",
    ],
    "start_date": ["Award Date", "Start Date (Manual)"],
    "mp_start_date": ["MP START", "MP Start Date", "MP Start Date (Manual)"],
    "report_date": ["Weekly Report Update Date (System)", "Report Date", "Update on"],
}
TIMELINES = {
    "first_crd_date": (["First CRD Date", "First CRD"], ["First CRD Date", "First CRD"]),
    "kick_off_date": (
        ["Eng , OEM Kick Off", "Kick Off Date"],
        ["Kick off", "Kick off date"],
    ),
    "tra": (["TRA"], ["TRA", "TRA/ECN DD", "ECN DD"]),
    "cut_steel": (["Cut Steel"], ["Cut Steel", "Cut Steel Date"]),
    "fot": (["FOT"], ["FOT", "FOT Date"]),
    "eb1": (["EB1"], ["EB1", "EB1 Date"]),
    "eb2": (["EB2"], ["EB2", "EB2 Date"]),
    "eb3": (["EB3"], ["EB3", "EB3 Date"]),
    "pilot": (["Pilot"], ["Pilot", "Pilot Date"]),
    "tooling_load": (
        ["Tooling Transfer Load"],
        ["Tooling Transfer Load", "Tooling Transfer Load Date"],
    ),
    "tooling_arrival": (
        ["Tooling Transfer Arrival"],
        ["Tooling Transfer Arrival", "Tooling Transfer Arrival Date"],
    ),
    "mpra": (["MPRA , ECN"], ["MPRA", "MPRA/ECN", "MPRA/ECN IMP"]),
    "mp_start_date": (
        ["MP START", "MP Start Date"],
        ["MP Start", "MP Start Date", "New MP Start", "New MP Start Date"],
    ),
    "previous_mp": (
        ["Previous MP Start Date"],
        [
            "Previous MP Date",
            "Previous MP Start",
            "Previous MP Start Date",
            "Old MP Start",
            "Old MP Start Date",
        ],
    ),
    "p1_cad_drop": (["P1 CAD DROP"], ["P1 CAD DROP", "P1 CAD"]),
    "p1_build_date": (["P1 BUILD DATE"], ["P1 BUILD", "P1 BUILD DATE"]),
    "p2_cad_drop": (["P2 CAD DROP"], ["P2 CAD DROP", "P2 CAD"]),
    "p2_build_date": (["P2 BUILD DATE"], ["P2 BUILD", "P2 BUILD DATE"]),
    "p3_cad_drop": (["P3 CAD DROP"], ["P3 CAD DROP", "P3 CAD"]),
    "p3_build_date": (["P3 BUILD DATE"], ["P3 BUILD", "P3 BUILD DATE"]),
    "last_p": (["Last P"], ["Last P"]),
    "last_p_build_date": (
        ["Last P BUILD DATE"],
        ["Last P BUILD", "Last P BUILD DATE"],
    ),
    "weeks_13_release": (
        ["13Weeks Release", "13 Weeks Release"],
        ["13Weeks", "13 Weeks", "13Weeks Release", "13 Weeks Release"],
    ),
    "as_release": (["AS Release"], ["AS Release", "AS"]),
    "color_approval": (["Color approval"], ["Color approval", "Color"]),
    "dqtp": (["DQTP"], ["DQTP", "DQTP Finish", "DQTP Finish Date"]),
    "compliance": (
        ["Compliance"],
        ["Compliance", "Compliance Complete", "Compliance Complete Date"],
    ),
    "fpo": (["FPO"], ["FPO", "FPO Release"]),
    "mp_aw_release": (
        ["MP AW release", "MP AW Release"],
        ["MP AW", "MP AW release", "MP AW Release"],
    ),
}


def fingerprint(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def scalar(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    # Excel stores time-only and elapsed-time cells separately from dates.
    # Preserve their precision without guessing a calendar date. The native
    # exporter still copies untouched cells from the original workbook XML.
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    return value


def read_tracker(path):
    import openpyxl

    path = Path(path).resolve()
    if path.suffix.lower() != ".xlsx" or not path.is_file():
        raise WorkbookError("请选择存在的 .xlsx 汇总表。")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise WorkbookError("汇总表超过 40 MB 上限。")
    digest = fingerprint(path)
    with ZipFile(path) as archive:
        if sum(p.file_size for p in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
            raise WorkbookError("汇总表解压后超过 200 MB 上限。")
        if len(archive.namelist()) != len(set(archive.namelist())):
            raise WorkbookError("汇总表含重复 ZIP 部件，无法安全更新。")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        wb = openpyxl.load_workbook(
            path, read_only=True, data_only=False, keep_links=False
        )
    try:
        sheets = [
            s
            for s in wb.worksheets
            if normalize_header(s.title) in {"allprojects", "alltracker", "alltrack"}
        ]
        if len(sheets) != 1:
            raise WorkbookError(
                "请提供含唯一 ALL PROJECTS（或 All Tracker）主表的汇总文件。"
            )
        sheet = sheets[0]
        effective_rows, effective_cols = sheet.max_row, sheet.max_column
        if (sheet.max_row or 0) * (sheet.max_column or 0) > 2_000_000:
            # Styled empty cells do not define the data range. Keep the original
            # ZIP intact; only bound this read by meaningful cells.
            import xml.etree.ElementTree as ET
            from openpyxl.utils.cell import coordinate_to_tuple
            effective_rows = effective_cols = 0
            with ZipFile(path) as archive, archive.open(sheet._worksheet_path) as stream:
                for _, element in ET.iterparse(stream, events=('end',)):
                    if element.tag.rsplit('}', 1)[-1] == 'c':
                        if any(child.tag.rsplit('}', 1)[-1] in {'v', 'f', 'is'} and ''.join(child.itertext())
                               for child in element):
                            r, c = coordinate_to_tuple(element.attrib['r'])
                            effective_rows, effective_cols = max(effective_rows, r), max(effective_cols, c)
                        element.clear()
            if effective_rows * effective_cols > 2_000_000:
                raise WorkbookError("总表有效数据范围过大，无法安全处理。")
        matrix = list(sheet.iter_rows(max_row=effective_rows, max_col=effective_cols))
        identity_labels = {normalize_header(x) for x in HEADERS["project_id"]}
        hits = [
            (i, c.column)
            for i, row in enumerate(matrix[:30], 1)
            for c in row
            if c.value is not None and normalize_header(c.value) in identity_labels
        ]
        if len(hits) != 1:
            raise WorkbookError(
                "前 30 行中未找到唯一的 Project Number / Project ID 表头。"
            )
        header_row, id_col = hits[0]
        headers = {
            c.column: clean_text(c.value)
            for c in matrix[header_row - 1]
            if c.value is not None
        }
        index = defaultdict(list)
        cells = {}
        formulas = set()
        notes = []
        date_headers = {
            normalize_header(alias)
            for key, aliases in HEADERS.items() if key.endswith("_date")
            for alias in aliases
        } | {
            normalize_header(alias)
            for aliases, _ in TIMELINES.values() for alias in aliases
        }
        for row in matrix[header_row:]:
            idcell = row[id_col - 1]
            if idcell.value and idcell.data_type != "f":
                index[normalize_project_id(idcell.value)].append(idcell.row)
            for cell in row:
                if cell.value is not None:
                    cells[cell.coordinate] = scalar(cell.value)
                    if (isinstance(cell.value, (time, timedelta))
                            and normalize_header(headers.get(cell.column, "")) in date_headers):
                        notes.append(
                            f"{sheet.title}!{cell.coordinate}（{headers[cell.column]}）"
                            f"含纯时间或时长 {scalar(cell.value)}，不是完整日期；"
                            "未推测日期，未更新的单元格将保留原值。"
                        )
                    if cell.data_type == "f":
                        formulas.add(cell.coordinate)
        if fingerprint(path) != digest:
            raise WorkbookError("读取期间汇总表发生变化，请重新生成预览。")
        content = {
            "headers": headers,
            "index": dict(index),
            "values": {k: v for k, v in cells.items() if k not in formulas},
        }
        content_hash = hashlib.sha256(
            json.dumps(content, sort_keys=True, ensure_ascii=False).encode()
        ).hexdigest()
        return {
            "path": str(path),
            "sha256": digest,
            "content_sha256": content_hash,
            "sheet": sheet.title,
            "sheet_part": sheet._worksheet_path,
            "headers": headers,
            "id_col": id_col,
            "header_row": header_row,
            "index": dict(index),
            "cells": cells,
            "formulas": formulas,
            "warnings": notes,
        }
    finally:
        wb.close()


def build_tracker_plan(path, reports):
    from openpyxl.utils import get_column_letter

    tracker = read_tracker(path)
    notes = list(tracker["warnings"])
    changes = []
    unmatched = []
    columns = {}
    versions = {}
    sidecar = Path(str(tracker["path"]) + ".autopm.json")
    if sidecar.exists():
        try:
            saved = json.loads(sidecar.read_text(encoding="utf-8"))
            if (
                saved.get("output_sha256") != tracker["sha256"]
                and saved.get("content_sha256") != tracker["content_sha256"]
                and saved.get('roundtrip', {}).get('version') != 1
            ):
                raise ValueError("changed")
            versions = saved["versions"]
            if not isinstance(versions, dict) or any(
                not isinstance(k, str) or not parse_date(v) for k, v in versions.items()
            ):
                raise ValueError("shape")
        except (ValueError, OSError, KeyError, TypeError):
            raise WorkbookError(
                "总表的同步记录损坏或与文件不一致，不能自动判断新旧版本。请核对该总表及旁边的 .autopm.json。"
            ) from None
    mappings = {
        **HEADERS,
        **{
            k: list(dict.fromkeys(HEADERS.get(k, []) + v[0]))
            for k, v in TIMELINES.items()
        },
    }
    for key, names in mappings.items():
        matches = [
            col
            for col, title in tracker["headers"].items()
            if normalize_header(title) in {normalize_header(x) for x in names}
        ]
        if len(matches) > 1:
            raise WorkbookError(f"总表存在多个可对应 {key} 的列，请核对重复表头。")
        if matches:
            columns[key] = matches[0]
    if "report_date" not in columns and not versions:
        notes.append(
            "总表无周报版本日期或历史同步记录；无法从 Date Added 判断新旧。导出会附带 .autopm.json，请与 Excel 一起保留。"
        )
    grouped = defaultdict(list)
    for report in reports:
        notes.extend(report.get("warnings", []))
        for p in report["projects"]:
            grouped[(normalize_project_id(p["project_id"]), p.get('_tracker_row'))].append((report, p))
    matched = 0
    source_projects = sum(len(r["projects"]) for r in reports)
    for (pid, matched_row), items in grouped.items():
        rows = [matched_row] if matched_row is not None and matched_row in tracker['index'].get(pid, []) else tracker["index"].get(pid, [])
        if len(rows) != 1:
            reason = f"{pid}：总表匹配到 {len(rows)} 行，已跳过项目。"
            notes.append(reason)
            unmatched.append(pid)
            continue
        dated = [
            (parse_date(p.get("report_date") or r.get("report_date")), r, p)
            for r, p in items
        ]
        if any(not d for d, _, _ in dated):
            notes.append(f"{pid}：周报日期无效，已跳过。")
            continue
        latest = max(d for d, _, _ in dated)
        selected = [(r, p) for d, r, p in dated if d == latest]
        # Equal-date conflicting reports are not ordered by filename or upload order.
        signatures = {
            json.dumps(
                {k: p.get(k) for k in ("fields", "tasks", "issues")},
                sort_keys=True,
                ensure_ascii=False,
            )
            for r, p in selected
        }
        if len(signatures) > 1:
            notes.append(f"{pid}：同日期的多份周报内容不同，已跳过项目。")
            continue
        report, p = selected[0]
        row = rows[0]
        matched += 1
        if len(dated) > len(selected):
            notes.append(f"{pid}：采用最新周报 {latest}，忽略较早版本。")
        version_key = p.get('_unit_key') or pid
        old_date = versions.get(version_key)
        if "report_date" in columns and p.get('date_basis') != 'airtable_capture':
            old = tracker["cells"].get(
                f"{get_column_letter(columns['report_date'])}{row}"
            )
            if old not in (None, ""):
                parsed = parse_date(old)
                if not parsed:
                    notes.append(f"{pid}：总表更新日期无法确定，已跳过项目。")
                    continue
                old_date = max(filter(None, [old_date, parsed]))
        if old_date and old_date > latest:
            notes.append(f"{pid}：周报 {latest} 早于总表版本 {old_date}，已跳过。")
            continue
        desired = defaultdict(list)

        def propose(key, value, is_date=False, source=None):
            if value in (None, "", []):
                return
            if key not in columns:
                notes.append(f"{pid}：{key} 在总表无明确对应列，保留在解析记录。")
                return
            if columns[key] in PROTECTED_TRACKER_COLUMNS:
                ref = f"{get_column_letter(columns[key])}{row}"
                notes.append(f"{pid}：{ref} 位于受保护的 P/Q/R 列，保留总表原值。")
                return
            if isinstance(value, (dict, list, bool)):
                notes.append(f"{pid}：{key} 值类型无法直接写入，已跳过。")
                return
            value = scalar(value)
            if is_date:
                value = parse_date(value)
                if not value:
                    notes.append(f"{pid}：{key} 日期不明确，已跳过。")
                    return
            if isinstance(value, str):
                value = clean_text(value)
                if len(value) > 32767:
                    notes.append(f"{pid}：{key} 超过 Excel 单元格长度上限，已跳过。")
                    return
            desired[columns[key]].append(
                (key, value, is_date, source or p.get("source", ""))
            )

        for key, value in p.get("fields", {}).items():
            if key != "project_id" and not (matched_row is not None and key in {'sku', 'factory', 'project_number'}):
                propose(key, value, key.endswith("_date"))
        aliases = {
            normalize_header(v): key
            for key, (_, names) in TIMELINES.items()
            for v in names
        }
        for task in p.get("tasks", []):
            key = aliases.get(normalize_header(task["name"]))
            if key:
                propose(key, task.get("date"), True, task.get("source", ""))
            else:
                notes.append(
                    f"{pid}：任务 {task['name']} 无唯一总表列（地区阶段不合并），已跳过。"
                )
        if p.get("issues"):
            notes.append(
                f"{pid}：{len(p['issues'])} 条问题/行动没有总表明细列，完整内容保留在同步记录。"
            )
        if "report_date" in columns and p.get('date_basis') != 'airtable_capture':
            propose("report_date", latest, True)
        for col, values in desired.items():
            ref = f"{get_column_letter(col)}{row}"
            before = tracker["cells"].get(ref)
            if (
                len({(json.dumps(v, ensure_ascii=False), d) for _, v, d, _ in values})
                != 1
            ):
                notes.append(f"{pid}：{ref} 有多个不同候选值，已跳过该单元格。")
                continue
            key, after, is_date, source = values[0]
            if ref in tracker["formulas"]:
                notes.append(f"{pid}：{ref} 是公式，保留原公式；周报候选值为 {after}。")
                continue
            if before == after:
                continue
            changes.append(
                {
                    "cell": ref,
                    "row": row,
                    "column": col,
                    "project_id": pid,
                    "field": tracker["headers"][col],
                    "semantic": key,
                    "before": before,
                    "after": after,
                    "is_date": is_date,
                    "source": source,
                    "report": report.get("source", ""),
                }
            )
        versions[version_key] = latest
    return {
        "version": 1,
        "target": "local_tracker",
        "tracker_path": tracker["path"],
        "tracker_sha256": tracker["sha256"],
        "sheet": tracker["sheet"],
        "sheet_part": tracker["sheet_part"],
        "changes": changes,
        "warnings": list(dict.fromkeys(notes)),
        "source_projects": source_projects,
        "matched_projects": matched,
        "unmatched_projects": unmatched,
        "versions": versions,
        "reports": reports,
        "source_files": {},
        "created_at": datetime.now().isoformat(timespec="seconds"),
    }


def verify_sources(plan):
    expected = {
        plan["tracker_path"]: plan["tracker_sha256"],
        **plan.get("source_files", {}),
    }
    for path, digest in expected.items():
        if not Path(path).is_file() or fingerprint(path) != digest:
            raise WorkbookError("汇总表或周报已变化/移动，请重新生成预览后导出。")


def export_tracker(plan, destination):
    """Preserve original native ZIP parts instead of round-tripping workbook objects."""
    from .native_xlsx import write_copy

    verify_sources(plan)
    destination = Path(destination).resolve()
    inputs = {
        Path(p).resolve() for p in [plan["tracker_path"], *plan.get("source_files", {})]
    }
    if destination in inputs:
        raise WorkbookError("不能覆盖输入汇总表或周报，请另存为新文件。")
    if destination.suffix.lower() != ".xlsx":
        raise WorkbookError("导出文件须为 .xlsx。")
    return write_copy(plan, destination)
