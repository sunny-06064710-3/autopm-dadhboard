"""One identity rule for cloud export and reviewed, three-way tracker writeback."""
from collections import defaultdict
from copy import deepcopy
import json
from pathlib import Path

from .identity import match_units, guard_matches
from .normalize import normalize_header, normalize_project_id, parse_date
from .schema_memory import DEFAULT_FIELDS
from .sync import (_field, _get, _mapping, _coerce, _equivalent, _resolve_factory,
                   _resolve_people, _finish_plan, DEPARTMENTS, WRITABLE)
from .tracker import (HEADERS, TIMELINES, PROTECTED_TRACKER_COLUMNS, read_tracker,
                      build_tracker_plan, fingerprint)
from .weekly_remark import uses_weekly_remark, read_weekly_remark, merge_weekly_remark
from .workbook import WorkbookError


def context(snapshot, config):
    ids = snapshot['table_ids']
    tables = {k: next(t for t in snapshot['schema']['tables'] if t['id'] == v) for k, v in ids.items()}
    fields = {k: {key: _field(tables[k], value) for key, value in _mapping(config, k, defaults).items()}
              for k, defaults in DEFAULT_FIELDS.items()}
    records = {k: snapshot['records'].get(v, []) for k, v in ids.items()}
    return ids, tables, fields, records


def columns(tracker):
    names = {**HEADERS, **{k: list(dict.fromkeys(HEADERS.get(k, []) + v[0])) for k, v in TIMELINES.items()}}
    result = {}
    for key, aliases in names.items():
        hits = [col for col, header in tracker['headers'].items() if normalize_header(header) in {normalize_header(a) for a in aliases}]
        if len(hits) > 1:
            raise WorkbookError(f'{key} 有多个表头，无法唯一映射。')
        if hits:
            result[key] = hits[0]
    return result


def cell_value(tracker, row, col):
    from openpyxl.utils import get_column_letter
    return tracker['cells'].get(f'{get_column_letter(col)}{row}')


def export_plan(path, snapshot, config, *, project_ids=None):
    from openpyxl.utils import get_column_letter
    from .airtable_tracker import snapshot_to_report
    tracker = read_tracker(path)
    cols = columns(tracker)
    if not {'project_id', 'sku', 'factory'} <= cols.keys():
        raise WorkbookError('总表必须含项目编号、SKU、Manufacturing Factory（或 Factory）三列。')
    ids, tables, fields, records = context(snapshot, config)
    identity_fields = {k: (fields['projects'].get(k) or {}).get('id') for k in ('project_id', 'project_number', 'sku', 'factory')}
    if not all(identity_fields.values()):
        raise WorkbookError('项目编号、SKU、工厂的云端映射不完整。')
    notes, selected = [], []
    scope = None if project_ids is None else {normalize_project_id(p) for p in project_ids}
    for pid, rows in tracker['index'].items():
        if scope is not None and pid not in scope:
            continue
        for row in rows:
            try:
                factory = _resolve_factory(cell_value(tracker, row, cols['factory']), tables['factories'], records['factories'],
                    config.get('factory_aliases', {}), config.get('field_mapping', {}).get('factories'))
                sku = cell_value(tracker, row, cols['sku'])
                matches, warnings = match_units(pid, sku, factory, records['projects'], identity_fields)
                notes.extend(f'总表第 {row} 行：{w}' for w in warnings)
                if warnings or len(matches) != 1:
                    if len(matches) > 1:
                        notes.append(f'{pid} 第 {row} 行对应多个 SKU 管理记录，不能把不同记录压成一行；请按 SKU＋工厂拆行。')
                    continue
                match = matches[0]
                selected.append({'row':row, 'number':pid, 'sku':sku, 'factory_ids':factory,
                                 'record_id':match['record']['id'], 'guard':match['guard']})
            except (ValueError, TypeError) as exc:
                notes.append(f'{pid} 第 {row} 行：{exc}；未匹配。')
    numbers_by_record = defaultdict(set)
    for item in selected:
        numbers_by_record[item['record_id']].add(item['number'])
    selected = [i for i in selected if len(numbers_by_record[i['record_id']]) == 1]
    for rid, numbers in numbers_by_record.items():
        if len(numbers) > 1:
            notes.append(f'{rid} 同一记录对应多个正式编号 {", ".join(sorted(numbers))}，已排除。')
    identities = {i['record_id']:i['number'] for i in selected}
    cloud = snapshot_to_report(snapshot, config, identities=identities)
    by_record = {p['airtable_record_id']:p for p in cloud['projects']}
    report = {**cloud, 'projects':[], 'warnings':cloud['warnings'] + notes}
    for item in selected:
        if item['record_id'] not in by_record:
            continue
        p = deepcopy(by_record[item['record_id']])
        p['_tracker_row'] = item['row']
        p['_unit_key'] = '|'.join((item['number'], ','.join(sorted(item['guard']['sku'].split(', '))), *item['factory_ids']))
        p['_identity'] = item
        report['projects'].append(p)
    plan = build_tracker_plan(path, [report])
    plan.update(source_kind='airtable', airtable_source=report['airtable_source'])
    changes = {c['cell']:c for c in plan['changes']}
    baseline = []
    for p in report['projects']:
        if plan['versions'].get(p['_unit_key']) != p['report_date']:
            continue
        row = p['_tracker_row']
        for key, col in cols.items():
            if key in {'project_id', 'project_number', 'sku', 'factory', 'report_date'}:
                continue
            if col in PROTECTED_TRACKER_COLUMNS:
                continue
            ref = f'{get_column_letter(col)}{row}'
            if ref in tracker['formulas']:
                continue
            source = p['field_sources'].get(key, '')
            source_parts = source.split('/')
            if len(source_parts) == 5 and source_parts[0] == 'Airtable':
                _, base, tid, rid, fid = source_parts
                kind = next((k for k, v in ids.items() if v == tid), None)
            elif fields['projects'].get(key) and key not in TIMELINES:
                kind, tid, rid, fid = 'projects', ids['projects'], p['airtable_record_id'], fields['projects'][key]['id']
            else:
                continue
            if kind not in {'projects', 'tasks'}:
                continue
            field = next((f for f in tables[kind]['fields'] if f['id'] == fid), None)
            record = next((r for r in records[kind] if r['id'] == rid), None)
            if not field or not record or field['type'] not in WRITABLE:
                continue
            after = changes[ref]['after'] if ref in changes else tracker['cells'].get(ref)
            if ref in changes:
                changes[ref]['source'] = source
            baseline.append({'identity':p['_identity'], 'header':tracker['headers'][col], 'semantic':key,
                             'kind':kind, 'table_id':tid, 'record_id':rid, 'field':field,
                             'cloud_before':_get(record, field), 'excel_before':after,
                             'report_date':p['report_date'],
                             'task_identity':({k:_get(record, fields['tasks'].get(k)) for k in ('project','name','milestone')} if kind == 'tasks' else None)})
    plan['roundtrip'] = {'version':1, 'base_id':snapshot['base_id'], 'bindings':baseline}
    if not baseline:
        plan['warnings'].append('没有可安全回写的单元格；请核对编号、SKU、工厂、周报日期及字段映射。')
    return plan


def writeback_plan(path, snapshot, config):
    path = Path(path).resolve()
    sidecar = Path(str(path) + '.autopm.json')
    if not sidecar.is_file():
        raise WorkbookError('缺少导出基线，请先通过 Airtable → All Tracker 导出，再编辑该副本。')
    source_hash, sidecar_hash = fingerprint(path), fingerprint(sidecar)
    saved = json.loads(sidecar.read_text(encoding='utf-8'))
    baseline = saved.get('roundtrip', {})
    if baseline.get('version') != 1 or baseline.get('base_id') != snapshot['base_id']:
        raise WorkbookError('该总表没有本数据库的回写基线，请重新从 Airtable 导出。')
    tracker = read_tracker(path)
    cols = columns(tracker)
    ids, tables, fields, records = context(snapshot, config)
    plan = {'version':1, 'base_id':snapshot['base_id'], 'source':str(path), 'source_kind':'tracker_writeback',
            'schema':snapshot['schema'], 'table_ids':ids, 'changes':[], 'project_guards':[],
            'warnings':list(tracker['warnings']), 'blockers':[],
            'source_files':{str(path):source_hash, str(sidecar):sidecar_hash}}
    pending, guards, verified_rows = {}, {}, {}
    for binding in baseline.get('bindings', []):
        identity = binding['identity']; pid = identity['number']; key = binding['semantic']
        try:
            unit = (identity['record_id'], pid, identity['sku'], tuple(identity['factory_ids']))
            if unit not in verified_rows:
                if not guard_matches(records['projects'], identity['guard']):
                    raise ValueError('编号＋SKU＋工厂已变化或不再唯一')
                rows = []
                for row in tracker['index'].get(pid, []):
                    if 'sku' not in cols or 'factory' not in cols:
                        break
                    from .identity import sku_tokens
                    if sku_tokens(cell_value(tracker, row, cols['sku'])) != sku_tokens(identity['sku']):
                        continue
                    linked = _resolve_factory(cell_value(tracker, row, cols['factory']), tables['factories'], records['factories'],
                        config.get('factory_aliases', {}), config.get('field_mapping', {}).get('factories'))
                    if sorted(linked) == sorted(identity['factory_ids']):
                        rows.append(row)
                if len(rows) != 1:
                    raise ValueError('编辑后总表中的编号＋SKU＋工厂不再唯一或已改动')
                verified_rows[unit] = rows[0]
            col = cols.get(key)
            if col is None:
                continue  # Removing a column is not deleting cloud data.
            if col in PROTECTED_TRACKER_COLUMNS:
                plan['warnings'].append(
                    f'{pid} / {binding.get("header", key)}：P/Q/R 为受保护列，未生成回写。'
                )
                continue
            from openpyxl.utils import get_column_letter
            ref = f'{get_column_letter(col)}{verified_rows[unit]}'
            after = cell_value(tracker, verified_rows[unit], col)
            if ref in tracker['formulas'] or after in (None, ''):
                continue
            field = next((f for f in tables[binding['kind']]['fields'] if f['id'] == binding['field']['id']), None)
            if not field or field['type'] != binding['field']['type'] or field.get('options') != binding['field'].get('options'):
                raise ValueError('云端字段类型或选项已变化，请重新导出')
            if _equivalent(after, binding['excel_before'], field):
                continue
            kind = binding['kind']; fid = field['id']; rid = binding['record_id']
            expected = fields[kind].get('due_date' if kind == 'tasks' else key)
            if not expected or expected['id'] != fid or ids[kind] != binding['table_id']:
                raise ValueError('字段或表映射已变化，请重新导出')
            if kind == 'projects' and rid != identity['record_id']:
                raise ValueError('回写基线的目标项目不一致')
            record = next((r for r in records[kind] if r['id'] == rid), None)
            if not record:
                raise ValueError('原目标记录已不存在')
            if kind == 'tasks':
                if any(_get(record, fields['tasks'].get(k)) != v for k, v in binding['task_identity'].items()):
                    raise ValueError('任务的项目、名称或里程碑已变化')
            before = _get(record, field)
            if key == 'current_progress' and uses_weekly_remark(config):
                value = merge_weekly_remark(binding['cloud_before'], binding['report_date'], {'current_progress':str(after)}, rich_text=field['type'] == 'richText')
            elif field['type'] == 'multipleRecordLinks':
                value = _resolve_people(after, tables['people'], records['people'], DEPARTMENTS.get(key),
                    config.get('people_aliases', {}), config.get('field_mapping', {}).get('people'))
            else:
                value = _coerce(field, after, binding['report_date'])
            if _equivalent(value, before, field):
                continue
            if not _equivalent(before, binding['cloud_before'], field):
                raise ValueError('云端与本地都修改了该字段，未覆盖云端')
            target = (binding['table_id'], rid)
            op = pending.setdefault(target, {'kind':kind, 'table_id':target[0], 'record_id':rid,
                'project_id':pid, 'project_record_id':identity['record_id'], 'report_date':binding['report_date'],
                'fields':{}, 'before':{}, 'field_names':{}, 'source':str(path)})
            if fid in op['fields'] and not _equivalent(op['fields'][fid], value, field):
                raise ValueError('多行对同一目标字段提出不同修改')
            op['fields'][fid] = value; op['before'][fid] = before; op['field_names'][fid] = field['name']
            if kind == 'tasks':
                from .normalize import normalize_key
                op['identity'] = {'project_field':fields['tasks']['project']['id'],
                                  'title_field':fields['tasks']['name']['id'],
                                  'title':normalize_key(binding['task_identity']['name'])}
                if fields['tasks'].get('milestone'):
                    op['matched_milestone'] = {'field':fields['tasks']['milestone']['id'], 'value':binding['task_identity']['milestone']}
            project = next(r for r in records['projects'] if r['id'] == identity['record_id'])
            version = fields['projects']['current_progress'] if uses_weekly_remark(config) else fields['projects']['report_date']
            raw_version = _get(project, version)
            guard = {'record_id':project['id'], 'project_id':pid, 'id_field':fields['projects']['project_id']['id'],
                     'date_field':version['id'], 'report_date':binding['report_date'], 'unit_identity':identity['guard']}
            if uses_weekly_remark(config):
                guard.update(date_storage='engineering_remark', remark_before=raw_version, remark_after=raw_version)
            guards[project['id']] = guard
        except (ValueError, TypeError, KeyError, StopIteration) as exc:
            plan['blockers'].append(f'{pid} / {binding.get("header", key)}：{exc}')
    for op in pending.values():
        guard = guards.get(op['project_record_id'])
        if guard and guard.get('date_storage') == 'engineering_remark' and op['kind'] == 'projects':
            guard['remark_after'] = op['fields'].get(guard['date_field'], guard['remark_before'])
    plan['changes'] = list(pending.values())
    plan['project_guards'] = list(guards.values())
    plan['blockers'] = list(dict.fromkeys(plan['blockers']))
    if fingerprint(path) != source_hash or fingerprint(sidecar) != sidecar_hash:
        raise WorkbookError('生成预览期间文件变化，请重新读取。')
    return _finish_plan(plan)
