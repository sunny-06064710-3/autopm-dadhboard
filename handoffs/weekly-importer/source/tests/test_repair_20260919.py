"""Regression checks for confirmed rules; no network or live records."""
import unittest
import tempfile
from pathlib import Path
from copy import deepcopy
from test_sync import fixture, FakeClient
from autopm.sync import build_plan, apply_plan, _resolve_factory
from autopm.identity import match_units
from autopm.tracker import TIMELINES
from autopm.weekly_remark import read_weekly_remark, merge_weekly_remark
from autopm.tracker_roundtrip import export_plan, writeback_plan
from autopm.tracker import export_tracker
import openpyxl


class ConfirmedRuleTests(unittest.TestCase):
    def test_review_shows_unit_and_complete_escaped_issue_content(self):
        from autopm.preview import enrich_display,save_preview
        snapshot,report=fixture()
        snapshot['records']['tbl_projects'][0]['fields'].update(fld_projects_sku='SKU-1',fld_projects_factory=['recFactory'])
        report['projects'][0]['fields'].update(sku='SKU-1',factory='F1')
        text='<img onerror=bad> '+'Detailed issue '*30
        report['projects'][0]['issues']=[{'text':text,'action':'Resolve & verify'}]
        plan=enrich_display(build_plan(report,snapshot,{'project_identity_mode':'number_sku_factory'}),snapshot)
        with tempfile.TemporaryDirectory() as folder:
            save_preview(report,plan,folder)
            html=(Path(folder)/'preview.html').read_text(encoding='utf-8')
        self.assertIn('SKU-1 / recFactory',html)
        self.assertIn(('Detailed issue '*30).rstrip(),html)
        self.assertIn('Resolve &amp; verify',html)
        self.assertNotIn('<img onerror=bad>',html)

    def test_factory_full_name_plus_stored_code_is_exact_not_fuzzy(self):
        snapshot,_=fixture()
        table=next(t for t in snapshot['schema']['tables'] if t['id']=='tbl_factories')
        records=snapshot['records']['tbl_factories']
        self.assertEqual(_resolve_factory('Factory One F1',table,records,{}),['recFactory'])
        with self.assertRaises(ValueError):_resolve_factory('Factory One F2',table,records,{})

    def test_old_multisection_remark_remains_readable_and_untouched(self):
        old = ('Manual note\n[[AUTOPM_WEEKLY_REPORT:v1]]\n周报日期: 2026-09-14\n'
               '工程进展:\nOld progress\n[[AUTOPM_WEEKLY_REPORT:END_PROGRESS]]\n'
               '本周更新:\nOld update\n[[AUTOPM_WEEKLY_REPORT:END_UPDATE]]\n'
               'Jira Open:\n2\n[[AUTOPM_WEEKLY_REPORT:END_JIRA_OPEN]]\n'
               '[[/AUTOPM_WEEKLY_REPORT]]')
        self.assertEqual(read_weekly_remark(old)['fields']['jira_open'], '2')
        new = merge_weekly_remark(old, '2026-09-19', {'current_progress': 'New progress'})
        self.assertTrue(new.startswith(old + '\n\n'))
        self.assertEqual(read_weekly_remark(new)['fields']['current_progress'], 'New progress')

    def test_line_endings_do_not_duplicate_plain_progress(self):
        stored = '--- 2026-09-19 ---\nLine one\nLine two'
        self.assertEqual(merge_weekly_remark(stored, '2026-09-19',
                         {'current_progress': 'Line one\r\nLine two\r\n'}), stored)

    def test_old_date_and_nested_history_are_rejected(self):
        stored = '--- 2026-09-19 ---\nCurrent'
        with self.assertRaises(ValueError):
            merge_weekly_remark(stored, '2026-09-18', {'current_progress': 'Old'})
        with self.assertRaises(ValueError):
            merge_weekly_remark(stored, '2026-09-19', {'current_progress': stored})

    def test_new_date_without_progress_change_advances_guard_without_copying_text(self):
        stored = '--- 2026-09-18 ---\nUnchanged progress'
        newer = merge_weekly_remark(stored,'2026-09-19',{})
        self.assertEqual(newer.count('Unchanged progress'),1)
        self.assertEqual(read_weekly_remark(newer), {'report_date':'2026-09-19','fields':{'current_progress':'Unchanged progress'}})
        with self.assertRaises(ValueError):
            merge_weekly_remark(newer,'2026-09-18',{'current_progress':'Old'})

    def test_text_current_stage_keeps_the_complete_weekly_update(self):
        snapshot, report = fixture()
        text = 'EB stage: tooling trial completed; next test on Friday'
        report['projects'][0]['fields']['update_this_week'] = text
        plan = build_plan(report, snapshot)
        change = next(c for c in plan['changes'] if c['kind'] == 'projects')
        self.assertEqual(change['fields']['fld_projects_update_this_week'], text)

    def test_last_p_is_distinct_from_last_p_build(self):
        self.assertEqual(TIMELINES['last_p'][0], ['Last P'])
        self.assertNotIn('Last P', TIMELINES['last_p_build_date'][1])

    def test_crd_fallback_explicit_priority_conflicts_and_no_source_mutation(self):
        for explicit, blocked, expected in [(None, [], '2026-09-17'), ('2026-09-20', [], '2026-09-20'),
                                            (None, ['first_crd_date'], None), (None, ['mp_start_date'], None)]:
            snapshot, report = fixture()
            p = report['projects'][0]
            p['fields'].update(mp_start_date='2026-11-01', first_crd_date=explicit)
            p['blocked_fields'] = blocked
            original = deepcopy(report)
            plan = build_plan(report, snapshot)
            tasks = [c for c in plan['changes'] if c['kind']=='tasks']
            if expected:
                self.assertEqual(tasks[0]['fields']['fld_tasks_due_date'],expected)
            else:
                self.assertFalse(tasks)
            self.assertEqual(report,original)

    def test_next_plm_inherits_npi_owner(self):
        snapshot, report = fixture()
        report['projects'][0]['fields']['next_plm'] = 'Release by Oct 15'
        op = next(c for c in build_plan(report,snapshot)['changes'] if c['kind']=='tasks')
        self.assertEqual(op['fields']['fld_tasks_owners'],['recNpi'])
        self.assertEqual(op['fields']['fld_tasks_due_date'],'2026-10-15')

    def test_merged_cloud_skus_are_not_updated_as_one_management_unit(self):
        snapshot,report=fixture()
        snapshot['records']['tbl_projects'][0]['fields'].update(fld_projects_sku='SKU-1, SKU-2',fld_projects_factory=['recFactory'])
        report['projects'][0]['fields'].update(sku='SKU-1',factory='F1')
        plan=build_plan(report,snapshot,{'project_identity_mode':'number_sku_factory'})
        self.assertFalse(plan['changes'])
        self.assertTrue(any('合并了多个 SKU' in w for w in plan['warnings']))

    def test_source_multiple_skus_fan_out_to_individual_records(self):
        snapshot,report=fixture()
        row=snapshot['records']['tbl_projects'][0]
        row['fields'].update(fld_projects_sku='SKU-1',fld_projects_factory=['recFactory'])
        other=deepcopy(row);other['id']='recSKU2';other['fields']['fld_projects_sku']='SKU-2'
        snapshot['records']['tbl_projects'].append(other)
        report['projects'][0]['fields'].update(sku='SKU-1, SKU-2',factory='F1')
        plan=build_plan(report,snapshot,{'project_identity_mode':'number_sku_factory'})
        self.assertEqual({c['record_id'] for c in plan['changes']},{'recProject','recSKU2'})

    def test_regional_milestone_conflicts_stop_and_same_dates_merge(self):
        for second, expected in [('2026-10-01',1),('2026-10-02',0)]:
            snapshot,report=fixture()
            snapshot['records']['tbl_projects'][0]['fields'].update(fld_projects_sku='SKU-1',fld_projects_factory=['recFactory'])
            report['projects'][0]['fields'].update(sku='SKU-1',factory='F1')
            report['projects'][0]['tasks']=[{'name':'EB1 (CN)','date':'2026-10-01'},{'name':'EB1 (VN)','date':second}]
            plan=build_plan(report,snapshot,{'project_identity_mode':'number_sku_factory'})
            tasks=[c for c in plan['changes'] if c['kind']=='tasks']
            self.assertEqual(len(tasks),expected)
            if tasks:self.assertEqual(tasks[0]['fields']['fld_tasks_name'],'EB1')

    def test_source_number_then_sku_factory_binds_x_record_without_renaming_id(self):
        snapshot, report = fixture()
        row = snapshot['records']['tbl_projects'][0]
        row['fields'].update(fld_projects_project_id='XLOCAL001', fld_projects_sku='SKU-1',
                             fld_projects_factory=['recFactory'])
        other = deepcopy(row); other['id'] = 'recOtherFactory'
        other['fields']['fld_projects_factory'] = ['recDifferentFactory']
        snapshot['records']['tbl_projects'].append(other)
        report['projects'][0]['fields'].update(sku='SKU-1', factory='F1')
        report['projects'][0]['tasks'] = [{'name':'EB1', 'date':'2026-10-01'}]
        plan = build_plan(report, snapshot, {'project_identity_mode':'number_sku_factory'})
        self.assertFalse(plan['blockers'], plan)
        self.assertEqual({c['project_record_id'] for c in plan['changes']}, {'recProject'})
        client = FakeClient(snapshot)
        with tempfile.TemporaryDirectory() as folder:
            result = apply_plan(client, plan, folder)
        self.assertEqual(result['status'], 'completed', result)
        fields = client.data['records']['tbl_projects'][0]['fields']
        self.assertEqual(fields['fld_projects_project_number'], 'AB-123')
        self.assertEqual(fields['fld_projects_project_id'], 'XLOCAL001')
        self.assertEqual(client.data['records']['tbl_tasks'][0]['fields']['fld_tasks_project'], ['recProject'])

    def test_same_sku_and_factory_with_other_official_number_is_not_reassigned(self):
        fields = {'project_id':'id','project_number':'number','sku':'sku','factory':'factory'}
        rows = [{'id':'rec1','fields':{'id':'X1','number':'AB-999','sku':'SKU-1','factory':['f1']}}]
        matches, notes = match_units('AB-123','SKU-1',['f1'],rows,fields)
        self.assertEqual(matches, [])
        self.assertTrue(notes)

    def test_duplicate_pair_and_missing_factory_are_not_guessed(self):
        fields = {'project_id':'id','project_number':'number','sku':'sku','factory':'factory'}
        rows = [{'id':'rec1','fields':{'id':'X1','sku':'SKU-1','factory':['f1']}},
                {'id':'rec2','fields':{'id':'X2','sku':'SKU-1','factory':['f1']}}]
        self.assertFalse(match_units('AB-123','SKU-1',['f1'],rows,fields)[0])
        self.assertFalse(match_units('AB-123','SKU-1',[],rows[:1],fields)[0])


class TrackerRoundtripTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name); self.source = self.root/'source.xlsx'; self.output = self.root/'output.xlsx'
        self.snapshot, _ = fixture()
        self.config = {'base_id':'appTest','project_identity_mode':'number_sku_factory'}
        for table in self.snapshot['schema']['tables']:
            for f in table['fields']:
                if f['id'] == 'fld_tasks_completed_by':
                    f.update(type='multipleRecordLinks', options={'linkedTableId':'tbl_people'})
        self.project = self.snapshot['records']['tbl_projects'][0]
        self.project['fields'].update(fld_projects_sku='SKU-1', fld_projects_factory=['recFactory'], fld_projects_status='On Track')
        self.snapshot['records']['tbl_tasks'] = [{'id':'recEB1','fields':{
            'fld_tasks_name':'EB1','fld_tasks_project':['recProject'],'fld_tasks_milestone':'EB1','fld_tasks_due_date':'2026-10-01'}}]
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'ALL PROJECTS'
        for c, v in enumerate(['Project Number','SKU','Manufacturing Factory','Project Name','Eng Status','EB1'],1):ws.cell(2,c,v)
        for c, v in enumerate(['AB-123','SKU-1','F1','Old name','On Track','2026-09-01'],1):ws.cell(3,c,v)
        wb.save(self.source); wb.close()

    def export(self):
        plan = export_plan(self.source, self.snapshot, self.config)
        self.assertTrue(plan['roundtrip']['bindings'], plan)
        export_tracker(plan, self.output)

    def edit(self, cell, value):
        wb = openpyxl.load_workbook(self.output); wb.active[cell] = value; wb.save(self.output); wb.close()

    def test_untouched_roundtrip_writes_nothing(self):
        self.export()
        plan = writeback_plan(self.output, self.snapshot, self.config)
        self.assertFalse(plan['blockers'], plan)
        self.assertEqual(plan['changes'], [])

    def test_protected_p_q_r_columns_are_never_exported_or_bound(self):
        self.snapshot['records']['tbl_tasks'].append({'id':'recKickOff','fields':{
            'fld_tasks_name':'Kick Off','fld_tasks_project':['recProject'],
            'fld_tasks_milestone':'Kick Off','fld_tasks_due_date':'2026-10-09'}})
        wb = openpyxl.load_workbook(self.source); ws = wb.active
        for col, heading, value in (
            (16, 'Eng , OEM Kick Off', '2026-09-01'),
            (17, 'Original TRA [at KO]', '2026-09-02'),
            (18, 'Original MP Ready [at KO]', '2026-09-03'),
        ):
            ws.cell(2, col, heading); ws.cell(3, col, value)
        wb.save(self.source); wb.close()
        plan = export_plan(self.source, self.snapshot, self.config)
        self.assertFalse(any(c['column'] in {'P', 'Q', 'R'} for c in plan['changes']))
        self.assertFalse(any(b['semantic'] == 'kick_off_date' for b in plan['roundtrip']['bindings']))
        self.assertTrue(any('受保护的 P/Q/R' in w for w in plan['warnings']))
        export_tracker(plan, self.output)
        self.edit('P3', '2026-10-10')
        reverse = writeback_plan(self.output, self.snapshot, self.config)
        self.assertFalse(reverse['changes'])


    def test_full_three_stage_chain_preserves_important_fields_and_is_repeatable(self):
        from test_rich_text_roundtrip import FormattingClient
        _, report = fixture()
        for table in self.snapshot['schema']['tables']:
            for field in table['fields']:
                if field['id'] == 'fld_projects_current_progress':
                    field['type'] = 'richText'
        self.config.update(base_id='appTest', project_report_storage='engineering_remark',project_report_storage_base_id='appTest')
        report['report_date']='2026-09-19'
        report['projects'][0]['fields'].update(sku='SKU-1',factory='F1',current_progress='Original progress\n3. Ship 10 units\n4. On Hold',
            update_this_week='EB testing complete; design review pending',loa_13weeks_plan='Complete 13-week plan',
            jira_total='0',jira_open='0')
        report['projects'][0]['tasks']=[{'name':'EB1','date':'2026-10-05'},{'name':'Custom engineering review','date':'2026-10-06'}]
        report['projects'][0]['issues']=[{'text':'Test failure requires review','action':'Repeat the test','owner':'Jane Smith','due_date':'2026-10-06','risk':'High'}]
        client=FormattingClient(self.snapshot)
        plan=build_plan(report,client.data,self.config)
        applied=apply_plan(client,plan,self.root/'weekly-import')
        self.assertEqual(applied['status'],'completed',applied)
        self.assertFalse(build_plan(report,client.data,self.config)['changes'])
        self.assertTrue(client.data['records']['tbl_issues'])
        wb=openpyxl.load_workbook(self.source)
        for col,heading in enumerate(['Current Progress','Eng Update This Week','LOA-13weeks plan','Jira summary'],7):wb.active.cell(2,col,heading)
        wb.save(self.source);wb.close()
        self.snapshot=client.data
        self.export()
        self.assertFalse(writeback_plan(self.output,client.data,self.config)['changes'])
        self.edit('G3','Corrected progress from tracker\n7. Ship 12 units\n8. Reviewed')
        reverse=writeback_plan(self.output,client.data,self.config)
        self.assertFalse(reverse['blockers'],reverse)
        result=apply_plan(client,reverse,self.root/'tracker-back')
        self.assertEqual(result['status'],'completed',result)
        self.assertFalse(writeback_plan(self.output,client.data,self.config)['changes'])
        cloud=client.data['records']['tbl_projects'][0]['fields']
        self.assertIn('Original progress',cloud['fld_projects_current_progress'])
        self.assertEqual(read_weekly_remark(cloud['fld_projects_current_progress'])['fields']['current_progress'],'Corrected progress from tracker\n1. Ship 12 units\n2. Reviewed')
        self.assertEqual(cloud['fld_projects_update_this_week'],'EB testing complete; design review pending')
        self.assertEqual(cloud['fld_projects_loa_13weeks_plan'],'Complete 13-week plan')
        self.assertEqual(cloud['fld_projects_jira_summary'],'Total: 0, Open: 0')

    def test_cloud_without_weekly_date_exports_with_capture_basis(self):
        self.snapshot['captured_at'] = '2026-09-19T05:00:00+00:00'
        self.project['fields'].pop('fld_projects_report_date',None)
        self.export()
        plan = writeback_plan(self.output,self.snapshot,self.config)
        self.assertFalse(plan['changes'])
        self.assertFalse(plan['blockers'])

    def test_zero_changed_cells_still_exports_writeback_baseline(self):
        self.export()
        plan = export_plan(self.output,self.snapshot,self.config)
        self.assertFalse(plan['changes'])
        other = self.root/'unchanged.xlsx'
        export_tracker(plan,other)
        self.assertTrue(other.is_file())
        self.assertTrue(plan['roundtrip']['bindings'])

    def test_mapping_change_and_new_duplicate_block_writeback(self):
        self.export(); self.edit('E3','At Risk')
        changed_config = deepcopy(self.config)
        changed_config.setdefault('field_mapping',{}).setdefault('projects',{})['status']='fld_projects_project_name'
        self.assertTrue(writeback_plan(self.output,self.snapshot,changed_config)['blockers'])
        plan = writeback_plan(self.output,self.snapshot,self.config)
        client = FakeClient(self.snapshot)
        duplicate = deepcopy(self.project); duplicate['id']='recDuplicate'
        client.data['records']['tbl_projects'].append(duplicate)
        result = apply_plan(client,plan,self.root/'duplicate')
        self.assertNotEqual(result['status'],'completed')
        self.assertFalse(client.writes)

    def test_changed_project_and_task_only_apply_to_original_records_and_repeat_is_empty(self):
        self.export(); self.edit('E3','At Risk'); self.edit('F3','2026-10-05')
        plan = writeback_plan(self.output, self.snapshot, self.config)
        self.assertFalse(plan['blockers'], plan)
        self.assertEqual({c['record_id'] for c in plan['changes']}, {'recProject','recEB1'})
        client = FakeClient(self.snapshot)
        result = apply_plan(client, plan, self.root/'run')
        self.assertEqual(result['status'], 'completed', result)
        self.assertFalse(writeback_plan(self.output, client.data, self.config)['changes'])

    def test_cloud_edit_conflict_blocks_entire_review(self):
        self.export(); self.edit('D3','Local edit')
        self.project['fields']['fld_projects_project_name'] = 'Cloud edit'
        plan = writeback_plan(self.output, self.snapshot, self.config)
        self.assertTrue(plan['blockers'])
        client = FakeClient(self.snapshot)
        result = apply_plan(client, plan, self.root/'conflict')
        self.assertEqual(result['status'], 'blocked'); self.assertFalse(client.writes)

    def test_blank_formula_and_missing_columns_do_not_clear_cloud(self):
        self.export(); self.edit('E3',None); self.edit('D3','=1+1')
        plan = writeback_plan(self.output, self.snapshot, self.config)
        self.assertFalse(plan['changes'])

    def test_wrong_base_and_changed_source_are_rejected(self):
        self.export(); self.edit('E3','At Risk')
        other = deepcopy(self.snapshot); other['base_id'] = 'appOther'
        with self.assertRaises(ValueError):writeback_plan(self.output,other,self.config)
        plan = writeback_plan(self.output,self.snapshot,self.config)
        self.edit('E3','On Track')
        with self.assertRaises(RuntimeError):apply_plan(FakeClient(self.snapshot),plan,self.root/'changed')

    def test_same_project_number_different_factory_rows_remain_separate(self):
        second = deepcopy(self.project); second['id'] = 'recSecond'; second['fields']['fld_projects_factory'] = ['recSecondFactory']
        self.snapshot['records']['tbl_projects'].append(second)
        self.snapshot['records']['tbl_factories'].append({'id':'recSecondFactory','fields':{'fld_factories_name':'Factory Two','fld_factories_code':'F2'}})
        wb = openpyxl.load_workbook(self.source)
        for c, v in enumerate(['AB-123','SKU-1','F2','Second old','On Track'],1):wb.active.cell(4,c,v)
        wb.save(self.source); wb.close()
        self.export(); self.edit('E4','At Risk')
        plan = writeback_plan(self.output,self.snapshot,self.config)
        self.assertFalse(plan['blockers'],plan)
        self.assertEqual([c['record_id'] for c in plan['changes']],['recSecond'])


if __name__ == '__main__':
    unittest.main()
